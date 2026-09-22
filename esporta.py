#!/usr/bin/env python3
"""
Esportazione dei dati scaricati (file JSON in output/) in tabelle leggibili.

Tabelle disponibili:
    insegnamenti   una riga per insegnamento di ogni piano (con n. scaglioni, docenti)
    scaglioni      una riga per scaglione (con docenti, moduli e sintesi dell'orario)
    lezioni        una riga per lezione settimanale (giorno, ora, aula, date)
    piani          una riga per corso/piano di studio (anche quelli scartati)

Formati: xlsx (Excel), csv (Excel italiano, separatore ';'), html (tabella con
ricerca e ordinamento), json, e "calendario" (orario settimanale in HTML).

Uso da terminale, esempi:
    python esporta.py output/manifesti.json --tabella scaglioni --formato xlsx
    python esporta.py output/manifesti.json --tabella lezioni --filtro corso_codice=531 --cerca analisi
    python esporta.py output/manifesti.json --calendario scaglione --filtro piano_codice=IT1
    python esporta.py output/manifesti.json --tutte --formato xlsx      (un foglio per tabella)
"""
import argparse
import colorsys
import csv
import hashlib
import html
import json
import re
import sys
from datetime import datetime
from pathlib import Path

HERE = Path(__file__).resolve().parent

# ------------------------------------------------------------------ colonne

LABELS = {
    "aa": "Anno accademico", "scuola": "Scuola", "corso_codice": "Cod. corso", "corso": "Corso di studi",
    "tipo_laurea": "Tipo laurea", "piano_codice": "Cod. piano", "piano": "Piano di studi",
    "piano_sede": "Sede piano", "piano_lingua": "Lingua piano", "anno_corso_piano": "Anno di corso (filtro)",
    "codice": "Codice", "insegnamento": "Insegnamento", "periodo": "Periodo", "cfu": "CFU",
    "tipo": "Tipo", "ssd": "SSD", "lingua": "Lingua", "sede_erogazione": "Sede erogazione",
    "blocco": "Blocco", "anno_corso": "Anno", "n_scaglioni": "N. scaglioni", "scaglioni": "Scaglioni",
    "docenti": "Docenti", "n_lezioni": "N. lezioni/sett.", "ore_settimanali": "Ore/sett.",
    "nota": "Note", "url": "Link dettaglio",
    "sezione": "Sezione", "scaglione": "Scaglione", "da": "Da (compreso)", "a": "A (escluso)",
    "moduli": "Moduli", "orario": "Orario", "aule": "Aule",
    "giorno": "Giorno", "giorno_n": "N. giorno", "inizio": "Inizio", "fine": "Fine",
    "durata_min": "Durata (min)", "aula": "Aula", "edificio": "Edificio / indirizzo",
    "attivita": "Attività", "dal": "Dal", "al": "Al", "n_date": "N. date", "date_lezioni": "Date lezioni",
    "semestre": "Semestre", "periodo_orario": "Periodo orario",
    "stato": "Stato", "n_insegnamenti_totali": "Insegnamenti nel piano", "n_insegnamenti": "Insegnamenti scaricati",
    "piani": "Piani", "corsi": "Corsi",
}

CONTESTO = ["aa", "scuola", "corso_codice", "corso", "tipo_laurea", "piano_codice", "piano",
            "piano_sede", "anno_corso_piano"]
# le colonne più utili vanno prima; il resto del contesto in fondo
_TESTA = ["corso", "piano_codice"]
_CODA = ["corso_codice", "piano", "piano_sede", "anno_corso_piano", "tipo_laurea", "scuola", "aa"]

COLONNE = {
    "insegnamenti": _TESTA + ["codice", "insegnamento", "periodo", "cfu", "n_scaglioni", "scaglioni", "docenti",
                              "n_lezioni", "tipo", "ssd", "lingua", "sede_erogazione", "blocco", "anno_corso",
                              "nota", "url"] + _CODA,
    "scaglioni": _TESTA + ["codice", "insegnamento", "periodo", "cfu", "scaglione", "docenti", "orario",
                           "ore_settimanali", "n_lezioni", "aule", "moduli", "da", "a", "sezione", "url"] + _CODA,
    "lezioni": _TESTA + ["codice", "insegnamento", "periodo", "scaglione", "giorno", "inizio", "fine", "aula",
                         "edificio", "docenti", "attivita", "dal", "al", "durata_min", "n_date", "date_lezioni",
                         "periodo_orario", "giorno_n", "da", "a", "url"] + _CODA,
    "piani": ["corso", "corso_codice", "anno_corso_piano", "piano_codice", "piano", "piano_sede", "piano_lingua",
              "stato", "n_insegnamenti_totali", "n_insegnamenti", "tipo_laurea", "scuola", "aa"],
}

DESCRIZIONI = {
    "insegnamenti": "Una riga per insegnamento di ogni piano: periodo, CFU, n. scaglioni, docenti",
    "scaglioni": "Una riga per scaglione: lettere da/a, docenti, moduli e sintesi dell'orario",
    "lezioni": "Una riga per lezione settimanale: giorno, ora, aula, date",
    "piani": "Una riga per corso e piano di studi (anche quelli scartati perché di altre sedi)",
}

# colonne che cambiano da un piano all'altro (la sezione è un codice interno del sito, diverso per ogni
# corso anche quando lo scaglione è lo stesso): ignorate quando si uniscono i duplicati
VARIABILI_PER_PIANO = set(CONTESTO) | {"url", "blocco", "anno_corso", "sezione"}

GIORNI_BREVI = {"Lunedì": "Lun", "Martedì": "Mar", "Mercoledì": "Mer", "Giovedì": "Gio",
                "Venerdì": "Ven", "Sabato": "Sab", "Domenica": "Dom"}


def label(col):
    return LABELS.get(col, col)


# ------------------------------------------------------------------ lettura

def carica(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def _lista(v):
    return ", ".join(v) if isinstance(v, list) else (v or "")


def _num(v):
    """'6.0' -> 6, '7,5' -> 7.5; il resto invariato."""
    try:
        f = float(str(v).replace(",", "."))
    except (TypeError, ValueError):
        return v
    return int(f) if f.is_integer() else f


# codice della bandierina sul sito -> lingua
LINGUE = {"it": "Italiano", "gb": "Inglese", "en": "Inglese", "uk": "Inglese", "fr": "Francese",
          "de": "Tedesco", "es": "Spagnolo", "pt": "Portoghese", "nl": "Olandese", "no": "Norvegese",
          "sv": "Svedese", "se": "Svedese", "pl": "Polacco", "zh": "Cinese", "cn": "Cinese"}


def _lingue(v):
    v = v if isinstance(v, list) else ([v] if v else [])
    return ", ".join(dict.fromkeys(LINGUE.get(str(x).lower(), str(x).upper()) for x in v))


def _nota_chiara(nota):
    """Note dello scraper in forma breve e leggibile (vale anche per i file scaricati
    con versioni precedenti, che contenevano il testo grezzo della pagina)."""
    if not nota:
        return ""
    if "Non esistono occupazioni" in nota or "Data Dove" in nota or nota == "nessun orario":
        return "nessuna lezione in orario"
    if nota == "tab orario disabilitato":
        return "orario non pubblicato sul sito"
    m = re.search(r"(Insegnamento riservato[^.]*?)\s+Erogato da\s+(.+?)\s+Città - Paese\s+(.+?)\s+Scheda", nota)
    if m:
        luogo = re.sub(r"^null\s*-\s*", "", m.group(3), flags=re.I).title()  # città mancante sul sito
        return f"{m.group(1)}; erogato da {m.group(2).title()} ({luogo})"
    return nota[:500]


def _ctx(dati, corso, piano):
    aa = (dati.get("meta", {}).get("parametri", {}) or {}).get("aa")
    anno = piano.get("anno_corso") if piano else None
    if anno is None and piano is not None:
        anno = (dati.get("meta", {}).get("parametri", {}) or {}).get("anno_corso")
    return {
        "aa": f"{aa}/{int(aa) + 1}" if aa and str(aa).isdigit() else aa,
        "scuola": corso["scuola"]["nome"], "corso_codice": corso["codice"], "corso": corso["nome"],
        "tipo_laurea": corso.get("tipo_ordinamento"),
        "piano_codice": (piano or {}).get("codice") or ("unico" if piano else None),
        "piano": (piano or {}).get("nome"), "piano_sede": (piano or {}).get("sede"),
        "anno_corso_piano": "tutti" if str(anno) == "0" else anno,
    }


def _nome_scaglione(da, a):
    nome = f"{da or '?'} – {a or '?'}"
    # A – ZZZZ copre tutti i cognomi: un solo scaglione per tutti gli studenti
    return f"{nome} (unico)" if da == "A" and a and set(a) == {"Z"} else nome


def _sintesi_orario(lezioni):
    parti = []
    for lz in sorted(lezioni, key=lambda x: (x.get("giorno_n") or 9, x.get("inizio") or "")):
        g = GIORNI_BREVI.get(lz.get("giorno"), lz.get("giorno") or "?")
        parti.append(f"{g} {lz.get('inizio')}-{lz.get('fine')} {lz.get('aula') or ''}".strip())
    return "; ".join(parti)


def tabelle(dati):
    """Dizionario {nome_tabella: lista di righe (dict)} dal JSON dello scraper."""
    out = {k: [] for k in COLONNE}
    for corso in dati.get("corsi_di_studio", []):
        for piano in corso.get("piani", []):
            ctx = _ctx(dati, corso, piano)
            out["piani"].append({**ctx, "piano_lingua": piano.get("lingua"), "stato": "incluso",
                                 "n_insegnamenti_totali": piano.get("n_insegnamenti_totali"),
                                 "n_insegnamenti": len(piano.get("insegnamenti", []))})
            for ins in piano.get("insegnamenti", []):
                base = {**ctx, "codice": ins.get("codice"), "insegnamento": ins.get("nome"),
                        "periodo": ins.get("periodo"), "cfu": _num(ins.get("cfu")), "url": ins.get("url_dettaglio")}
                sezioni = ins.get("sezioni") or []
                tutte_lez, tutti_doc, nomi_sc = [], [], []
                for sez in sezioni:
                    lez_sez = sez.get("orario") or []
                    tutte_lez += lez_sez
                    for sc in sez.get("scaglioni", []):
                        nome_sc = _nome_scaglione(sc.get("da"), sc.get("a"))
                        nomi_sc.append(nome_sc)
                        tutti_doc += [d for d in sc.get("docenti", []) if d not in tutti_doc]
                        lez_sc = [lz for lz in lez_sez
                                  if (lz.get("scaglione_da"), lz.get("scaglione_a")) == (sc.get("da"), sc.get("a"))]
                        moduli = [r.get("modulo") for r in sc.get("righe", []) if r.get("modulo")]
                        out["scaglioni"].append({
                            **base, "sezione": sez.get("id_sezione"), "scaglione": nome_sc,
                            "da": sc.get("da"), "a": sc.get("a"), "docenti": _lista(sc.get("docenti")),
                            "moduli": "; ".join(dict.fromkeys(moduli)), "n_lezioni": len(lez_sc),
                            "ore_settimanali": _num(round(sum(lz.get("durata_min") or 0 for lz in lez_sc) / 60, 2)),
                            "orario": _sintesi_orario(lez_sc) or _nota_chiara(sez.get("orario_nota")),
                            "aule": ", ".join(dict.fromkeys(lz.get("aula") for lz in lez_sc if lz.get("aula"))),
                        })
                    doc_sc = {(sc.get("da"), sc.get("a")): _lista(sc.get("docenti")) for sc in sez.get("scaglioni", [])}
                    for lz in lez_sez:
                        date = lz.get("date_lezioni") or []
                        out["lezioni"].append({
                            **base, "scaglione": _nome_scaglione(lz.get("scaglione_da"), lz.get("scaglione_a")),
                            "da": lz.get("scaglione_da"), "a": lz.get("scaglione_a"),
                            "docenti": doc_sc.get((lz.get("scaglione_da"), lz.get("scaglione_a")), ""),
                            "giorno": lz.get("giorno"), "giorno_n": lz.get("giorno_n"),
                            "inizio": lz.get("inizio"), "fine": lz.get("fine"), "durata_min": lz.get("durata_min"),
                            "aula": lz.get("aula"), "edificio": lz.get("aula_descrizione"),
                            "attivita": lz.get("attivita"), "dal": lz.get("dal"), "al": lz.get("al"),
                            "n_date": len(date) or None, "date_lezioni": ", ".join(date),
                            "periodo_orario": lz.get("periodo_orario"),
                        })
                if ins.get("errore"):
                    nota = f"non letto per un errore: {ins['errore']}"[:500]
                else:
                    nota = _nota_chiara(ins.get("nota_dettaglio")) or "; ".join(
                        dict.fromkeys(_nota_chiara(s["orario_nota"]) for s in sezioni if s.get("orario_nota")))
                out["insegnamenti"].append({
                    **base, "tipo": ins.get("tipo"), "ssd": ins.get("ssd"), "lingua": _lingue(ins.get("lingua")),
                    "sede_erogazione": ins.get("sede_erogazione"), "blocco": ins.get("blocco"),
                    "anno_corso": _num(ins.get("anno_corso")), "n_scaglioni": ins.get("n_scaglioni"),
                    "scaglioni": "; ".join(nomi_sc), "docenti": ", ".join(tutti_doc),
                    "n_lezioni": len(tutte_lez), "nota": nota,
                })
        for ps in corso.get("piani_scartati", []):
            ctx = _ctx(dati, corso, {"codice": ps.get("codice"), "nome": ps.get("nome"),
                                     "sede": ps.get("sede"), "anno_corso": ps.get("anno_corso")})
            out["piani"].append({**ctx, "stato": "scartato (altra sede)"})
        if not corso.get("piani") and not corso.get("piani_scartati"):
            ctx = _ctx(dati, corso, None)
            stato = (f"non letto per un errore: {corso['errore']}" if corso.get("errore")
                     else "; ".join(corso.get("note", [])) or corso.get("nota") or "nessun piano")
            out["piani"].append({**ctx, "stato": stato})
    return out


# ------------------------------------------------------------------ filtri

def valori_distinti(righe, col):
    vals = {str(r.get(col)) for r in righe if r.get(col) not in (None, "")}
    return sorted(vals, key=lambda v: (not v[:1].isdigit(), v.lower()))


def filtra(righe, filtri=None, cerca=""):
    """filtri: {colonna: insieme di valori ammessi (confronto come stringa)}.
    cerca: testo cercato (senza maiuscole) in qualunque colonna; più parole = tutte presenti."""
    filtri = {k: {str(x) for x in v} for k, v in (filtri or {}).items() if v}
    parole = [p for p in (cerca or "").lower().split() if p]
    out = []
    for r in righe:
        if any(str(r.get(k)) not in v for k, v in filtri.items()):
            continue
        if parole:
            testo = " ".join(str(x) for x in r.values() if x is not None).lower()
            if not all(p in testo for p in parole):
                continue
        out.append(r)
    return out


def unisci_duplicati(righe, colonne):
    """Unisce le righe identiche a parte il piano di studi (lo stesso insegnamento
    compare in più piani): i campi del piano diventano elenchi separati da ' | '."""
    chiave_cols = [c for c in colonne if c not in VARIABILI_PER_PIANO]
    gruppi = {}
    for r in righe:
        k = tuple(str(r.get(c)) for c in chiave_cols)
        if k not in gruppi:
            gruppi[k] = {c: [] for c in colonne}
        for c in colonne:
            v = r.get(c)
            if v not in (None, "") and v not in gruppi[k][c]:
                gruppi[k][c].append(v)
    out = []
    for vals in gruppi.values():
        riga = {}
        for c, v in vals.items():
            riga[c] = v[0] if len(v) == 1 else (" | ".join(str(x) for x in v) if v else None)
        out.append(riga)
    return out


# ------------------------------------------------------------------ scrittura tabelle

def _csv_val(v):
    if v is None:
        return ""
    if isinstance(v, float):
        return str(v).replace(".", ",")  # Excel italiano: '7.5' diventerebbe una data
    return v


def esporta_csv(path, colonne, righe):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow([label(c) for c in colonne])
        for r in righe:
            w.writerow([_csv_val(r.get(c)) for c in colonne])


def esporta_json(path, colonne, righe):
    Path(path).write_text(json.dumps([{c: r.get(c) for c in colonne} for r in righe],
                                     ensure_ascii=False, indent=1), encoding="utf-8")


def esporta_xlsx(path, fogli):
    """fogli: {nome_foglio: (colonne, righe)}"""
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    def val(v):
        return ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v

    wb = Workbook()
    wb.remove(wb.active)
    for nome, (colonne, righe) in fogli.items():
        ws = wb.create_sheet(nome[:31])
        ws.append([label(c) for c in colonne])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E79")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for r in righe:
            ws.append([val(r.get(c)) for c in colonne])
        if "url" in colonne:
            ci = colonne.index("url") + 1
            for row in ws.iter_rows(min_row=2, min_col=ci, max_col=ci):
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("http") and " | " not in cell.value:
                        cell.hyperlink = cell.value
                        cell.value = "apri"
                        cell.font = Font(color="0563C1", underline="single")
        for i, c in enumerate(colonne, 1):
            lung = max([len(label(c))] + [len(str(r.get(c) or "")) for r in righe[:500]])
            ws.column_dimensions[get_column_letter(i)].width = min(max(lung + 2, 8), 60)
        ws.freeze_panes = "A2"
        if righe:
            ws.auto_filter.ref = ws.dimensions
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


CSS_BASE = """
:root{color-scheme:light;--bg:#f7f8fa;--card:#fff;--fg:#1c1e21;--muted:#65676b;--line:#dde1e6;--head:#1f4e79;--headfg:#fff;--hi:#fff6d6}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--fg);font:14px/1.4 system-ui,-apple-system,"Segoe UI",Roboto,sans-serif}
header{padding:16px}h1{font-size:20px;margin:0 0 4px}.sub{color:var(--muted);font-size:13px}
"""


def esporta_html(path, colonne, righe, titolo, sottotitolo=""):
    dati = json.dumps({"cols": [label(c) for c in colonne],
                       "rows": [[r.get(c) for c in colonne] for r in righe]}, ensure_ascii=False)
    dati = dati.replace("</", "<\\/")
    pagina = f"""<!doctype html><html lang="it"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>{html.escape(titolo)}</title>
<style>{CSS_BASE}
.bar{{display:flex;gap:8px;flex-wrap:wrap;align-items:center;padding:0 16px 12px}}
input{{font:inherit;padding:7px 10px;border:1px solid var(--line);border-radius:6px;background:var(--card);color:var(--fg);min-width:260px;flex:1;max-width:480px}}
.wrap{{overflow:auto;margin:0 16px 16px;border:1px solid var(--line);border-radius:8px;background:var(--card);max-height:calc(100vh - 140px)}}
table{{border-collapse:collapse;width:max-content;min-width:100%}}
th{{position:sticky;top:0;background:var(--head);color:var(--headfg);text-align:left;padding:8px;cursor:pointer;white-space:nowrap;user-select:none}}
td{{padding:6px 8px;border-top:1px solid var(--line);vertical-align:top;max-width:340px}}
tr:hover td{{background:var(--hi)}} a{{color:#2f7fd8}}
</style></head><body>
<header><h1>{html.escape(titolo)}</h1><div class="sub">{html.escape(sottotitolo)}</div></header>
<div class="bar"><input id="q" placeholder="Cerca in tutte le colonne…" autofocus><span class="sub" id="n"></span>
<span class="sub">Clic su un'intestazione per ordinare</span></div>
<div class="wrap"><table><thead><tr id="h"></tr></thead><tbody id="b"></tbody></table></div>
<script>
const D={dati};let rows=D.rows.slice(),sortC=-1,asc=true;
const esc=s=>String(s).replace(/[&<>"]/g,c=>({{'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}}[c]));
const cell=v=>v==null?'':(typeof v==='string'&&v.startsWith('http')&&!v.includes(' | ')?'<a href="'+esc(v)+'" target="_blank">apri</a>':esc(v));
document.getElementById('h').innerHTML=D.cols.map((c,i)=>'<th data-i="'+i+'">'+esc(c)+'</th>').join('');
function draw(){{const q=document.getElementById('q').value.toLowerCase().split(/\\s+/).filter(Boolean);
 let r=rows.filter(x=>{{const t=x.join(' ').toLowerCase();return q.every(p=>t.includes(p))}});
 document.getElementById('n').textContent=r.length+' righe su '+D.rows.length;
 document.getElementById('b').innerHTML=r.slice(0,5000).map(x=>'<tr>'+x.map(v=>'<td>'+cell(v)+'</td>').join('')+'</tr>').join('');}}
document.getElementById('h').onclick=e=>{{const i=+e.target.dataset.i;if(isNaN(i))return;asc=sortC===i?!asc:true;sortC=i;
 rows.sort((a,b)=>{{const x=a[i],y=b[i];if(x==null)return 1;if(y==null)return -1;
 const r=(typeof x==='number'&&typeof y==='number')?x-y:String(x).localeCompare(String(y),'it',{{numeric:true}});return asc?r:-r}});draw();}};
document.getElementById('q').oninput=draw;draw();
</script></body></html>"""
    Path(path).write_text(pagina, encoding="utf-8")


# ------------------------------------------------------------------ calendario

RAGGRUPPA = {
    "scaglione": "Un orario per ogni insegnamento e scaglione",
    "insegnamento": "Un orario per insegnamento (tutti gli scaglioni insieme)",
    "piano": "Un orario per piano di studi (tutte le lezioni del piano)",
    "aula": "Un orario per aula",
    "docente": "Un orario per docente",
    "tutto": "Un unico orario con tutte le lezioni filtrate",
}


def _colore(testo):
    h = int(hashlib.md5(str(testo).encode()).hexdigest()[:6], 16) / 0xFFFFFF
    r, g, b = colorsys.hls_to_rgb(h, 0.86, 0.55)
    r2, g2, b2 = colorsys.hls_to_rgb(h, 0.35, 0.6)
    return f"#{int(r * 255):02x}{int(g * 255):02x}{int(b * 255):02x}", f"#{int(r2 * 255):02x}{int(g2 * 255):02x}{int(b2 * 255):02x}"


def _min(hhmm):
    h, m = str(hhmm).split(":")
    return int(h) * 60 + int(m)


def _gruppi_calendario(lezioni, per):
    gruppi = {}
    for lz in lezioni:
        if per == "scaglione":
            keys = [f"{lz.get('codice')} {lz.get('insegnamento')} · scaglione {lz.get('scaglione')}"]
        elif per == "insegnamento":
            keys = [f"{lz.get('codice')} {lz.get('insegnamento')}"]
        elif per == "piano":
            keys = [f"{lz.get('corso')} · piano {lz.get('piano_codice')}"]
        elif per == "aula":
            keys = [f"Aula {lz.get('aula') or '?'}"]
        elif per == "docente":
            keys = [d.strip() for d in str(lz.get("docenti") or "?").replace(" | ", ", ").split(",") if d.strip()]
        else:
            keys = ["Tutte le lezioni"]
        for k in keys:
            gruppi.setdefault(k, []).append(lz)
    return gruppi


def esporta_calendario(path, lezioni, per="scaglione", titolo="Orario settimanale", sottotitolo=""):
    lezioni = [lz for lz in lezioni if lz.get("inizio") and lz.get("fine") and lz.get("giorno_n")]
    gruppi = _gruppi_calendario(lezioni, per)
    parti = []
    for nome in sorted(gruppi, key=str.lower):
        lz_g = gruppi[nome]
        # stessa lezione presente in più piani -> una sola
        uniche = {}
        for lz in lz_g:
            k = (lz.get("codice"), lz.get("scaglione"), lz.get("giorno_n"), lz.get("inizio"), lz.get("aula"))
            uniche.setdefault(k, lz)
        lz_g = sorted(uniche.values(), key=lambda x: (x["giorno_n"], x["inizio"]))
        giorni = list(range(1, max(5, max(int(x["giorno_n"]) for x in lz_g)) + 1))
        t0 = min(_min(x["inizio"]) for x in lz_g) // 60 * 60
        t1 = -(-max(_min(x["fine"]) for x in lz_g) // 60) * 60
        # corsie per le lezioni sovrapposte nello stesso giorno
        corsie = {g: [] for g in giorni}
        pos = []
        for x in lz_g:
            g, a, b = int(x["giorno_n"]), _min(x["inizio"]), _min(x["fine"])
            for i, fine in enumerate(corsie[g]):
                if fine <= a:
                    corsie[g][i] = b
                    break
            else:
                i = len(corsie[g])
                corsie[g].append(b)
            pos.append((x, g, i, a, b))
        n_corsie = {g: max(1, len(corsie[g])) for g in giorni}
        col0 = {}
        c = 2
        for g in giorni:
            col0[g] = c
            c += n_corsie[g]
        tmpl = "56px " + " ".join(f"repeat({n_corsie[g]},minmax(0,1fr))" for g in giorni)
        n_righe = (t1 - t0) // 15
        celle = []
        nomi_giorni = ["Lunedì", "Martedì", "Mercoledì", "Giovedì", "Venerdì", "Sabato", "Domenica"]
        for g in giorni:
            celle.append(f'<div class="dh" style="grid-column:{col0[g]}/span {n_corsie[g]};grid-row:1">{nomi_giorni[g - 1]}</div>')
            celle.append(f'<div class="dcol" style="grid-column:{col0[g]}/span {n_corsie[g]};grid-row:2/span {n_righe}"></div>')
        for h in range(t0, t1, 60):
            r = (h - t0) // 15 + 2
            celle.append(f'<div class="hr" style="grid-row:{r}/span 4">{h // 60:02d}:00</div>')
            celle.append(f'<div class="hl" style="grid-column:2/-1;grid-row:{r}"></div>')
        for x, g, i, a, b in pos:
            bg, fg = _colore(x.get("codice"))
            date = f"dal {x.get('dal')} al {x.get('al')}" if x.get("dal") else ""
            tip = "\n".join(filter(None, [f"{x.get('codice')} {x.get('insegnamento')}", f"Scaglione {x.get('scaglione')}",
                                          x.get("attivita") if x.get("attivita") != x.get("insegnamento") else "",
                                          f"{x.get('giorno')} {x.get('inizio')}-{x.get('fine')}",
                                          f"Aula {x.get('aula')} – {x.get('edificio') or ''}", x.get("docenti"), date]))
            testo = [f"<b>{html.escape(str(x.get('insegnamento') or ''))}</b>"]
            if per != "scaglione":
                testo.append(f"sc. {html.escape(str(x.get('scaglione')))}")
            testo.append(f"{x['inizio']}–{x['fine']} · {html.escape(str(x.get('aula') or ''))}")
            celle.append(
                f'<div class="ev" title="{html.escape(tip)}" style="grid-column:{col0[g] + i};'
                f'grid-row:{(a - t0) // 15 + 2}/{(b - t0) // 15 + 2};background:{bg};border-color:{fg};color:#1c1e21">'
                + "<br>".join(testo) + "</div>")
        info = sorted({f"{x.get('corso_codice')}/{x.get('piano_codice')}" for x in lz_g if x.get("piano_codice")})
        periodi = sorted({f"{x.get('dal')} → {x.get('al')}" for x in lz_g if x.get("dal")})
        parti.append(
            f'<section><h2>{html.escape(nome)}</h2><div class="sub">{len(lz_g)} lezioni settimanali'
            + (f" · piani {html.escape(', '.join(info[:8]))}{'…' if len(info) > 8 else ''}" if per != "piano" and info else "")
            + (f" · periodo {html.escape(periodi[0])}" if len(periodi) == 1 else "")
            + f'</div><div class="cal" style="grid-template-columns:{tmpl};grid-template-rows:28px repeat({n_righe},11px)">'
            + "".join(celle) + "</div></section>")
    indice = "".join(f'<a href="#s{i}">{html.escape(n)}</a>' for i, n in enumerate(sorted(gruppi, key=str.lower)))
    corpo = "".join(p.replace("<section>", f'<section id="s{i}">', 1) for i, p in enumerate(parti))
    pagina = f"""<!doctype html><html lang="it"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>{html.escape(titolo)}</title>
<style>{CSS_BASE}
nav{{padding:0 16px 8px;display:flex;flex-wrap:wrap;gap:4px 12px;font-size:12px;max-height:160px;overflow:auto}}
nav a{{color:var(--muted)}} section{{margin:12px 16px 24px;background:var(--card);border:1px solid var(--line);border-radius:10px;padding:12px;overflow-x:auto}}
h2{{font-size:16px;margin:0 0 2px}} .cal{{display:grid;min-width:640px;margin-top:10px;position:relative}}
.dh{{text-align:center;font-weight:600;font-size:13px;border-bottom:2px solid var(--line)}}
.dcol{{border-left:1px solid var(--line)}} .hr{{grid-column:1;font-size:11px;color:var(--muted);transform:translateY(-7px)}}
.ev b{{display:block;white-space:nowrap;overflow:hidden;text-overflow:ellipsis}}
.hl{{border-top:1px dashed var(--line);pointer-events:none}}
.ev{{border-left:4px solid;border-radius:5px;margin:1px 2px;padding:2px 5px;font-size:11px;line-height:1.25;overflow:hidden;z-index:1}}
@media print{{nav{{display:none}} section{{break-inside:avoid}}}}
</style></head><body>
<header><h1>{html.escape(titolo)}</h1><div class="sub">{html.escape(sottotitolo)} · {len(gruppi)} orari · passa il mouse su una lezione per i dettagli</div></header>
<nav>{indice}</nav>{corpo or '<p style="padding:16px">Nessuna lezione con orario nei dati filtrati.</p>'}</body></html>"""
    Path(path).write_text(pagina, encoding="utf-8")
    return len(gruppi)


# ------------------------------------------------------------------ esporta (unica entrata)

FORMATI = {"xlsx": "Excel (.xlsx)", "csv": "CSV per Excel (.csv)", "html": "Pagina web con ricerca (.html)",
           "json": "JSON (.json)"}


def esporta(path, formato, nome_tabella, colonne, righe, sottotitolo=""):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if formato == "xlsx":
        esporta_xlsx(path, {nome_tabella: (colonne, righe)})
    elif formato == "csv":
        esporta_csv(path, colonne, righe)
    elif formato == "html":
        esporta_html(path, colonne, righe, f"PoliMi – {nome_tabella}", sottotitolo)
    elif formato == "json":
        esporta_json(path, colonne, righe)
    else:
        raise ValueError(f"formato sconosciuto: {formato}")
    return path


def descrivi_file(dati):
    p = dati.get("meta", {}).get("parametri", {}) or {}
    parti = [f"a.a. {p.get('aa')}/{int(p['aa']) + 1}" if str(p.get("aa", "")).isdigit() else "",
             p.get("sede_nome") or p.get("sede") or "",
             f"generato il {dati.get('meta', {}).get('generato_il', '?').replace('T', ' ')}"]
    meta = dati.get("meta", {})
    if meta.get("completo") is False:
        motivo = "per un errore" if meta.get("errore") else "interrotto" if meta.get("interrotto") else ""
        parti.append(f"SCARICAMENTO INCOMPLETO {motivo}".strip())
    return " · ".join(x for x in parti if x)


def main():
    ap = argparse.ArgumentParser(description="Esporta i dati scaricati in Excel/CSV/HTML/JSON o calendario",
                                 formatter_class=argparse.RawDescriptionHelpFormatter, epilog=__doc__)
    ap.add_argument("file", help="file JSON prodotto dallo scraper")
    ap.add_argument("--tabella", choices=list(COLONNE), default="insegnamenti")
    ap.add_argument("--tutte", action="store_true", help="tutte le tabelle (xlsx: un foglio ciascuna)")
    ap.add_argument("--formato", choices=list(FORMATI), default="xlsx")
    ap.add_argument("--filtro", action="append", default=[], metavar="COLONNA=VAL1,VAL2",
                    help="tieni solo le righe con quei valori (ripetibile). Colonne: " + ", ".join(COLONNE["lezioni"]))
    ap.add_argument("--cerca", default="", help="parole da cercare in qualunque colonna")
    ap.add_argument("--colonne", help="colonne da esportare, separate da virgola (default: tutte)")
    ap.add_argument("--unisci-duplicati", action="store_true",
                    help="una sola riga per insegnamento, scaglione o lezione anche se compare in più corsi e piani")
    ap.add_argument("--calendario", choices=list(RAGGRUPPA), help="crea l'orario settimanale HTML raggruppato così")
    ap.add_argument("--out", help="file di destinazione (default: in output/esportazioni/)")
    ap.add_argument("--elenca-valori", metavar="COLONNA", help="mostra i valori presenti in una colonna ed esce")
    a = ap.parse_args()
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    dati = carica(a.file)
    tab = tabelle(dati)
    filtri = {}
    for f in a.filtro:
        k, _, v = f.partition("=")
        filtri[k.strip()] = [x.strip() for x in v.split(",")]
    if a.elenca_valori:
        for v in valori_distinti(tab[a.tabella], a.elenca_valori):
            print(v)
        return
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    cartella = HERE / "output" / "esportazioni"

    def prepara(nome):
        cols = [c for c in (a.colonne.split(",") if a.colonne else COLONNE[nome]) if c in COLONNE[nome]]
        righe = filtra(tab[nome], {k: v for k, v in filtri.items() if k in COLONNE[nome]}, a.cerca)
        if a.unisci_duplicati:
            righe = unisci_duplicati(righe, cols)
        return cols, righe

    if a.calendario:
        # righe complete: il calendario ha bisogno di giorno/ora e toglie da sé i doppioni
        righe = filtra(tab["lezioni"], {k: v for k, v in filtri.items() if k in COLONNE["lezioni"]}, a.cerca)
        out = Path(a.out or cartella / f"calendario_{a.calendario}_{stamp}.html")
        out.parent.mkdir(parents=True, exist_ok=True)
        n = esporta_calendario(out, righe, a.calendario, sottotitolo=descrivi_file(dati))
        print(f"{n} orari salvati in {out}")
        return
    if a.tutte:
        if a.formato != "xlsx":
            sys.exit("--tutte è disponibile solo con --formato xlsx")
        out = Path(a.out or cartella / f"manifesti_tutte_{stamp}.xlsx")
        esporta_xlsx(out, {n: prepara(n) for n in COLONNE})
        print(f"Salvato in {out}")
        return
    cols, righe = prepara(a.tabella)
    out = Path(a.out or cartella / f"{a.tabella}_{stamp}.{a.formato}")
    esporta(out, a.formato, a.tabella, cols, righe, descrivi_file(dati))
    print(f"{len(righe)} righe salvate in {out}")


if __name__ == "__main__":
    main()
